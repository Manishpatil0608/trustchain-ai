# -*- coding: utf-8 -*-
"""
build_and_export_presentation.py
Master compiler that builds the upgraded 18-slide PowerPoint presentation,
validates shapes and notes, and synchronizes deliverables to docs/presentation/final/.
"""

import os
import shutil
import sys
from pptx import Presentation

from build_upgraded_presentation import prs, PPTX_PATH, ASSETS_DIR
import upgraded_slides_1_to_6
import upgraded_slides_7_to_12
import upgraded_slides_13_to_18

print("Building Upgraded TrustChain-AI Capstone Presentation...")

upgraded_slides_1_to_6.build_slides_1_to_6(prs, ASSETS_DIR)
upgraded_slides_7_to_12.build_slides_7_to_12(prs, ASSETS_DIR)
upgraded_slides_13_to_18.build_slides_13_to_18(prs, ASSETS_DIR)

# Save master PPTX
prs.save(PPTX_PATH)
print("PPTX successfully saved to:", PPTX_PATH)
print("File size:", os.path.getsize(PPTX_PATH), "bytes")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
# Final destination paths
FINAL_DIR = os.path.join(BASE_DIR, "docs", "presentation", "final")
FINAL_PPTX = os.path.join(FINAL_DIR, "TrustChain_AI_Final_Presentation.pptx")
ROOT_PPTX_AKTA = os.path.join(BASE_DIR, "TrustChain_AI_Final_Presentation.pptx")

os.makedirs(FINAL_DIR, exist_ok=True)
shutil.copy2(PPTX_PATH, FINAL_PPTX)
shutil.copy2(PPTX_PATH, ROOT_PPTX_AKTA)
print("Copied PPTX to final folder and root directory.")

# Validation check
saved_prs = Presentation(PPTX_PATH)
print(f"Total slides in presentation: {len(saved_prs.slides)}")
for i, slide in enumerate(saved_prs.slides):
    notes_chars = len(slide.notes_slide.notes_text_frame.text) if slide.has_notes_slide else 0
    print(f"Slide {i+1:02d}: shapes={len(slide.shapes):02d}, notes_chars={notes_chars}")

print("PPTX Build & Validation completed successfully.")
