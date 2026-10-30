# -*- coding: utf-8 -*-
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from build_full_deck import (
    COLOR_BG, COLOR_PANEL, COLOR_PANEL_ALT, COLOR_BORDER, COLOR_CYAN,
    COLOR_GREEN, COLOR_AMBER, COLOR_RED, COLOR_PURPLE, COLOR_WHITE,
    COLOR_SLATE, COLOR_MUTED, FONT_TITLE, FONT_BODY, FONT_MONO,
    create_base_slide, add_card, add_textbox, add_para, attach_notes
)

def build_slides_13_to_18(prs, assets_dir):
    # ==========================================================================
    # SLIDE 13: Multi-Agent Collaboration Sessions & Sequential Pipeline
    # ==========================================================================
    slide13 = create_base_slide("COLLABORATION ENGINE \u00B7 DETERMINISTIC PIPELINE", "Multi-Agent Collaboration Sessions & Sequential Pipeline Orchestration", 13)

    # Left Column: Bilateral Collaboration Protocol
    c13_l = add_card(slide13, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "Bilateral Collaboration Protocol & Lifecycle Management")
    tf13_l = add_textbox(slide13, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf13_l, "1. Session Lifecycle State Machine", size=11, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf13_l, "Collaboration sessions progress through deterministic states:\n  [ PENDING ] \u2192 [ ACTIVE ] \u2192 [ COMPLETED ] / [ TERMINATED ] / [ FAILED ]\n\u2022 PENDING: Session initiated, mutual identity and VC exchange verified.\n\u2022 ACTIVE: Bilateral task payload execution and sensor telemetry logging.\n\u2022 COMPLETED: Task successfully delivered, mutual ratings recorded.\n\u2022 TERMINATED: Forced termination upon trust decay or policy violation.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf13_l, "2. Session Entity Data Model (akta/vdr/models.py)", size=11, bold=True, color=COLOR_CYAN, mono=True, space_before=8)
    add_para(tf13_l, "\u2022 session_id: UUIDv4 unique session handle.\n\u2022 initiator_did & target_did: Cryptographic did:key identifiers.\n\u2022 session_token: Cryptographically signed ephemeral JWT.\n\u2022 collaboration_events: Granular append-only event log capturing every request, intermediate step, response payload, and sensor signature.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf13_l, "3. Mid-Session Trust Boundary Enforcement", size=11, bold=True, color=COLOR_AMBER, mono=True, space_before=8)
    add_para(tf13_l, "If either agent's trust drops below 0.50 during a multi-turn collaboration, the session is instantly aborted, preventing tainted output propagation.", size=9.5, color=COLOR_WHITE, space_before=2)

    # Right Column: Sequential Pipeline Orchestration
    c13_r = add_card(slide13, Inches(6.883), Inches(1.65), Inches(5.65), Inches(5.25), "Sequential Workflow Pipeline & Deterministic Simulation Engine")
    tf13_r = add_textbox(slide13, Inches(7.133), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf13_r, "SEQUENTIAL PIPELINE EXECUTION MODEL", size=11, bold=True, color=COLOR_PURPLE, mono=True)
    add_para(tf13_r, "Complex enterprise tasks are decomposed into multi-stage pipelines:\n  [ Node 1: Ingestion ] \u2192 [ Node 2: Analysis ] \u2192 [ Node 3: Verification ]", size=10, color=COLOR_WHITE, mono=True, space_before=2)
    add_para(tf13_r, "\u2022 Ordered Node Traversal: Output of Node i serves as verified input to Node i+1.\n\u2022 Adaptive Authorization at Every Step: Before each node executes, the gateway verifies the agent's real-time trust state (T \u2265 0.80 for privileged steps).\n\u2022 Atomic Failure Handling: If any node fails or is quarantined, the pipeline halts cleanly with full audit trail logging.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf13_r, "LOCAL AGENT CONNECTOR (REPRODUCIBILITY)", size=11, bold=True, color=COLOR_GREEN, mono=True, space_before=8)
    add_para(tf13_r, "\u2022 Deterministic Simulation: Implemented in akta/connectors/local.py.\n\u2022 Eliminates Live LLM Dependencies: Avoids flaky external API costs, network jitter, and token limits during automated test runs.\n\u2022 Realistic Workload Emulation: Simulates parameterized response payloads, synthetic compute delays, and injection error codes.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf13_r, "PARALLEL EXECUTION BOUNDARY (ACADEMIC HONESTY)", size=11, bold=True, color=COLOR_AMBER, mono=True, space_before=8)
    add_para(tf13_r, "While DAG dependencies are supported in the database schema (workflow_nodes.depends_on), current execution is strictly sequential.", size=9.5, color=COLOR_SLATE, space_before=2)

    attach_notes(slide13, {
        'narrative': "Slide 13 details our Multi-Agent Collaboration Engine and Pipeline Orchestration. When two agents collaborate, they form a bilateral session managed by a formal state machine: PENDING, ACTIVE, and either COMPLETED or TERMINATED. The session entity binds the initiator DID, target DID, task payload, and an ephemeral session token. In addition to bilateral sessions, we support sequential multi-stage pipelines where complex tasks flow through Node 1, Node 2, and Node 3. Crucially, before every single node is executed, our Adaptive Policy Enforcer re-verifies the agent's live trust state. If an agent degrades mid-pipeline, execution is immediately halted, stopping cascading corruption in its tracks. For local testing, our LocalAgentConnector simulates realistic agent execution deterministically without external LLM API costs.",
        'takeaway': "Collaboration sessions and sequential workflows enforce real-time trust gating at every step, halting instantly if an agent degrades mid-pipeline.",
        'evidence': "akta/collaboration_router.py, akta/vdr/models.py (CollaborationSession, Workflow, WorkflowNode), akta/connectors/local.py, tests/test_collaboration.py",
        'qa': "Q1: Can two agents collaborate if one is in SUSPICIOUS state?\nA1: Only for standard read/query tasks. Privileged workflow execution and delegation require both agents to be in TRUSTED state (T >= 0.80).\n\nQ2: Does your workflow engine support parallel branches?\nA2: The database schema includes a depends_on field for DAG parallel dependencies, but runtime execution in the capstone is strictly sequential to guarantee deterministic audit verification.",
        'caveat': "Workflow execution is tested deterministically with LocalAgentConnector to guarantee 100% reproducible test verification."
    })

    # ==========================================================================
    # SLIDE 14: Local Relational VDR Architecture & Quad-Dimensional Auditability
    # ==========================================================================
    slide14 = create_base_slide("STORAGE LAYER \u00B7 ZERO TAMPERING EVIDENCE", "Local Relational VDR Architecture & Quad-Dimensional Auditability", 14)

    # Left Column: Relational VDR Schema
    c14_l = add_card(slide14, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "13-Table Relational Schema & WAL High-Throughput Storage")
    tf14_l = add_textbox(slide14, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf14_l, "1. Verifiable Data Registry (VDR) Implementation", size=11, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf14_l, "TrustChain-AI implements a local relational VDR using SQLite configured with Write-Ahead Logging (WAL mode):\n\u2022 Sub-5ms Read/Write: High-throughput local operations without gas fees, block confirmation delays, or external RPC network failures.\n\u2022 ACID Compliance: Atomic transaction guarantees for all trust state updates and credential issuance operations.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf14_l, "2. Complete 13-Table Relational Architecture", size=11, bold=True, color=COLOR_CYAN, mono=True, space_before=8)
    add_para(tf14_l, "\u2022 Identity & Credentials: stored_vcs, agents\n\u2022 Trust & Telemetry: agent_trust_states, telemetry_logs, trust_audit_trails\n\u2022 Collaboration: collaboration_sessions, collaboration_events, agent_ratings\n\u2022 Workflows: workflows, workflow_nodes\n\u2022 Security & Access: users, refresh_tokens, security_audit_events", size=9.5, color=COLOR_WHITE, space_before=2)

    add_para(tf14_l, "3. Academic Clarification: Relational VDR vs Blockchain", size=11, bold=True, color=COLOR_AMBER, mono=True, space_before=8)
    add_para(tf14_l, "Per W3C DID specifications, a Verifiable Data Registry is a logical interface. Our implementation uses SQLite WAL for high performance. Deploying to an Ethereum/Hedera ledger represents future distributed work.", size=9.5, color=COLOR_SLATE, space_before=2)

    # Right Column: Quad-Dimensional Auditability
    c14_r = add_card(slide14, Inches(6.883), Inches(1.65), Inches(5.65), Inches(5.25), "Quad-Dimensional Auditability & Verifiable Ledger Traceability")
    tf14_r = add_textbox(slide14, Inches(7.133), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf14_r, "FOUR IMMUTABLE AUDIT STREAMS", size=11, bold=True, color=COLOR_PURPLE, mono=True)
    add_para(tf14_r, "TrustChain-AI maintains 4 distinct, tamper-evident audit streams providing 360-degree observability:", size=10, color=COLOR_SLATE, space_before=2)

    add_para(tf14_r, "[DIMENSION 1] stored_vcs (Credential Lifecycle)", size=10, bold=True, color=COLOR_CYAN, mono=True, space_before=6)
    add_para(tf14_r, "Records every issued, validated, or revoked W3C Verifiable Credential with cryptographic issuer signature and LD-Proof payload.", size=9, color=COLOR_SLATE, space_before=1)

    add_para(tf14_r, "[DIMENSION 2] trust_audit_trails (State Transitions)", size=10, bold=True, color=COLOR_GREEN, mono=True, space_before=6)
    add_para(tf14_r, "Immutable log of all trust updates: old_score, new_score, old_state, new_state, trigger_reason, and millisecond timestamp.", size=9, color=COLOR_SLATE, space_before=1)

    add_para(tf14_r, "[DIMENSION 3] telemetry_logs (Runtime Interaction Events)", size=10, bold=True, color=COLOR_AMBER, mono=True, space_before=6)
    add_para(tf14_r, "Granular per-task sensor telemetry: success, policy compliance, SLA latency, input hashes, and sensor signatures.", size=9, color=COLOR_SLATE, space_before=1)

    add_para(tf14_r, "[DIMENSION 4] security_audit_events (Security Gateway)", size=10, bold=True, color=COLOR_RED, mono=True, space_before=6)
    add_para(tf14_r, "Authentication attempts, account lockouts, RBAC privilege violations, and adaptive policy rejection intercepts.", size=9, color=COLOR_SLATE, space_before=1)

    attach_notes(slide14, {
        'narrative': "Slide 14 presents our storage architecture and quad-dimensional auditability. In the W3C specification, a Verifiable Data Registry is a logical concept, not necessarily a distributed blockchain. For TrustChain-AI, we engineered a high-throughput relational VDR using SQLite in WAL mode. This guarantees atomic transactions and sub-5 millisecond response times without gas costs or consensus delays. The database consists of exactly 13 tables partitioned across five functional domains. To guarantee forensic accountability, we maintain four separate audit dimensions: stored_vcs for credential provenance, trust_audit_trails for state transitions, telemetry_logs for raw interaction metrics, and security_audit_events for access control violations. Every event is timestamped and immutable.",
        'takeaway': "A high-performance 13-table SQLite WAL registry provides sub-5ms latency and quad-dimensional forensic auditability across credentials, trust, telemetry, and security.",
        'evidence': "akta/vdr/models.py, akta/vdr/database.py, akta/vdr/registry_crud.py, tests/test_platform_features.py",
        'qa': "Q1: Why didn't you use an actual blockchain like Ethereum or Polygon?\nA1: Blockchain introduces 2-15 second confirmation latencies and monetary gas costs per transaction. For autonomous agent interactions occurring at millisecond scale, a local WAL-enabled VDR provides the required throughput while strictly adhering to W3C VDR interface principles.\n\nQ2: How do you prevent tampering with the SQLite audit logs?\nA2: The backend enforces append-only semantics at the ORM layer, and each telemetry event includes cryptographic signatures tied to the agent's did:key.",
        'caveat': "The registry runs locally; transition to a Byzantine fault tolerant distributed ledger is identified as future work."
    })

    # ==========================================================================
    # SLIDE 15: Operational SOC Portal Demonstration & Interface Verification
    # ==========================================================================
    slide15 = create_base_slide("FRONTEND IMPLEMENTATION \u00B7 SECURITY OPERATIONS CENTER", "Operational SOC Portal Demonstration & Interface Verification", 15)

    # Left Column: Enterprise SOC Dashboard Architecture
    c15_l = add_card(slide15, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "Enterprise SOC Dashboard & Navigation Architecture")
    tf15_l = add_textbox(slide15, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf15_l, "1. Modern Frontend Architecture", size=11, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf15_l, "\u2022 React 19 + Vite: Ultra-responsive component architecture compiled to production in 26.2s via Vite.\n\u2022 Dark Navy SOC Theme: Styled with Tailwind CSS (#080E1A background, #111C35 panels, #00F0FF accents) with seamless light mode toggle.\n\u2022 Lucide Icons & Headless UI: Accessible, keyboard-navigable operational controls.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf15_l, "2. Five Dedicated Operational Modules", size=11, bold=True, color=COLOR_CYAN, mono=True, space_before=8)
    add_para(tf15_l, "\u2022 Overview / Metrics Dashboard: High-level KPI summary, trust distribution pie/bar charts, active collaboration count, live security alert feeds.\n\u2022 Agent Explorer (/agents): Grid view of all registered agents with real-time status badges (TRUSTED, SUSPICIOUS, QUARANTINED, REVOKED).\n\u2022 Discovery Center (/discovery): Capability tag selector, dynamic trust threshold slider (0.0\u20131.0), and diagnostic eligibility badges.\n\u2022 Matching Center (/matching): Task requirement input, multi-factor ranking, top candidate spotlight card, and explainable disqualification accordion.\n\u2022 User Management (/admin/users): Role-based access control (ADMIN, OPERATOR, VIEWER) with customized badges and audit logs.", size=9, color=COLOR_SLATE, space_before=2)

    # Right Column: Investigation Drawer & Explainability
    c15_r = add_card(slide15, Inches(6.883), Inches(1.65), Inches(5.65), Inches(5.25), "Interactive Investigation Drawer & Explainable Decision Displays")
    tf15_r = add_textbox(slide15, Inches(7.133), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf15_r, "AGENT INVESTIGATION SLIDE-OVER DRAWER", size=11, bold=True, color=COLOR_GREEN, mono=True)
    add_para(tf15_r, "Clicking 'INVESTIGATE' on any agent card opens a comprehensive slide-over drawer with 5 tabbed analytical views:", size=10, color=COLOR_WHITE, space_before=2)

    add_para(tf15_r, "[TAB 1] Overview: Agent DID, display name, verified capabilities, uptime, and current trust badge.", size=9.5, color=COLOR_SLATE, space_before=4)
    add_para(tf15_r, "[TAB 2] Trust History: Visual progression of historical EMA scores, degradation events, and recovery milestones.", size=9.5, color=COLOR_SLATE, space_before=3)
    add_para(tf15_r, "[TAB 3] Telemetry Stream: Real-time table of recent interactions (success rate, policy compliance %, SLA latency ms).", size=9.5, color=COLOR_SLATE, space_before=3)
    add_para(tf15_r, "[TAB 4] Verifiable Credentials: W3C JSON-LD credential viewer, cryptographic proof validity, and authorized capabilities.", size=9.5, color=COLOR_SLATE, space_before=3)
    add_para(tf15_r, "[TAB 5] Actions: Privileged operator controls for manual trust adjustment, temporary quarantine, or permanent revocation.", size=9.5, color=COLOR_SLATE, space_before=3)

    add_para(tf15_r, "EXPLAINABLE DECISION INTERFACE", size=11, bold=True, color=COLOR_PURPLE, mono=True, space_before=8)
    add_para(tf15_r, "Every matching decision displays explicit mathematical breakdown chips (Trust: 40%, Rep: 30%, Success: 20%, Latency: 10%) so human operators never face black-box recommendations.", size=9.5, color=COLOR_WHITE, space_before=2)

    attach_notes(slide15, {
        'narrative': "Slide 15 showcases our Security Operations Center portal. Built with React 19 and Vite, the interface provides a unified, professional cockpit for monitoring autonomous agents. The design follows enterprise SOC conventions with a dark navy palette, high-contrast cyan accents, and complete light mode compatibility. The portal features five dedicated views: the Executive Dashboard, the Agent Explorer, the Discovery Center, the Matching Center, and User Management. When an operator clicks 'Investigate' on any agent, a five-tab slide-over drawer appears, allowing deep inspection into the agent's DID, historical trust trajectory, raw telemetry events, and Verifiable Credentials. Furthermore, our matching interface provides full explainability chips, detailing exact mathematical score contributions.",
        'takeaway': "The React 19 SOC portal delivers comprehensive operational visibility, featuring a 5-tab investigation slide-over drawer and fully explainable decision interfaces.",
        'evidence': "frontend/src/App.jsx, frontend/src/components/AgentDrawer.jsx, frontend/src/components/DiscoveryCenter.jsx, frontend/src/components/MatchingCenter.jsx, frontend/src/components/UserManagement.jsx",
        'qa': "Q1: Does the frontend make direct security decisions?\nA1: No. The frontend is strictly an operational presentation layer. All trust evaluations, matching scores, and RBAC rules are enforced authoritatively on the FastAPI backend.\n\nQ2: How is the investigation drawer invoked?\nA2: The AgentCard component dispatches an onInvestigate callback passing the agent's DID, which loads the drawer asynchronously via GET /api/v1/registry/agents/{did}/details.",
        'caveat': "The UI is fully connected to the FastAPI backend and verified with clean production builds."
    })

    # ==========================================================================
    # SLIDE 16: Empirical Testing, Subsystem Health & Verification
    # ==========================================================================
    slide16 = create_base_slide("EMPIRICAL VALIDATION \u00B7 80/80 AUTOMATED TESTS PASSING", "Empirical Testing, Subsystem Health & Mathematical Verification", 16)

    # Left Column: Pytest Suite Summary
    c16_l = add_card(slide16, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "Automated Pytest Suite & 100% Verification Pass Rate")
    tf16_l = add_textbox(slide16, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf16_l, "VERIFIED AUTOMATED TEST SUITE: 80 PASSED IN 19.10s", size=11, bold=True, color=COLOR_GREEN, mono=True)
    add_para(tf16_l, "Every core framework component is validated by rigorous automated pytest test suites with zero failures and zero regressions:", size=9.5, color=COLOR_SLATE, space_before=2)

    test_suites = [
        ("test_crypto.py (8 tests)", "Ed25519 key generation, did:key multicodec 0xED01, base58-btc, signature verification"),
        ("test_vc.py (12 tests)", "W3C VC schema validation, URDNA2015 canonicalization, LD-Proofs, tampering detection"),
        ("test_trust_engine.py (14 tests)", "Telemetry ingestion, sliding window, EMA decay/recovery, 4 discrete state transitions"),
        ("test_adaptive_auth.py (10 tests)", "RBAC + ABAC enforcement, privileged vs standard action gating, HTTP 403 intercepts"),
        ("test_collaboration.py (11 tests)", "Bilateral session lifecycle, sequential workflow pipelines, mid-session quarantine"),
        ("test_platform_features.py (9 tests)", "13-table SQLite VDR operations, multi-factor matching formula, peer ratings"),
        ("test_investigation_discovery_matching.py (10 tests)", "Investigation drawer endpoints, capability discovery filters, exclusion rationales"),
        ("test_auth_rbac.py (6 tests)", "Argon2id password hashing, JWT refresh rotation, 5-strike lockout, role matrices")
    ]
    for name, desc in test_suites:
        add_para(tf16_l, f"\u2713 {name}", size=9.5, bold=True, color=COLOR_CYAN, space_before=3)
        add_para(tf16_l, f"   {desc}", size=8.5, color=COLOR_SLATE, space_before=1)

    add_para(tf16_l, "Frontend Build: Vite production bundle built in 26.20s with 0 errors.", size=9.5, color=COLOR_WHITE, space_before=6)

    # Right Column: Chart Image & Health Metrics
    chart3_img = os.path.join(assets_dir, "test_results_chart.png")
    if os.path.exists(chart3_img):
        slide16.shapes.add_picture(chart3_img, Inches(6.883), Inches(1.65), Inches(5.65), Inches(3.6))
    else:
        c16_r = add_card(slide16, Inches(6.883), Inches(1.65), Inches(5.65), Inches(3.6), "Automated Test Distribution Chart")

    # Health Card below chart
    c16_sub = add_card(slide16, Inches(6.883), Inches(5.4), Inches(5.65), Inches(1.5), "Live Subsystem Health Probing (/health)", fill_color=COLOR_PANEL_ALT)
    tf16_sub = add_textbox(slide16, Inches(7.133), Inches(5.8), Inches(5.15), Inches(1.0))
    add_para(tf16_sub, "HEALTH PROBE RESPONSE (GET /health):", size=9, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf16_sub, '{\n  "status": "healthy",  "vdr": "sqlite_wal_connected",\n  "active_agents": 8,  "pass_rate": "100% (80/80 tests)"\n}', size=9.5, color=COLOR_GREEN, mono=True, space_before=2)

    attach_notes(slide16, {
        'narrative': "Slide 16 details our empirical testing and subsystem verification. In computer science capstone projects, implementation claims must be backed by reproducible evidence. We verified the complete framework against an 80-test automated pytest suite, achieving 100% pass rate in 19.10 seconds across eight distinct test modules. These tests cover cryptographic primitives, Verifiable Credential tampering, mathematical EMA boundary conditions, adaptive policy enforcement, collaboration lifecycles, and RBAC security. The chart on the upper right shows the distribution across the eight suites. On the frontend, Vite compiles the production bundle in 26.2 seconds with zero syntax or packaging errors. Finally, our live /health endpoint continuously confirms that the SQLite WAL database and active agents are fully operational.",
        'takeaway': "The complete framework is rigorously validated by 80 automated unit and integration tests passing in 19.10 seconds with 0 failures.",
        'evidence': "tests/ (8 test modules, 80 test functions), pytest.ini, frontend/package.json (npm run build), akta/server.py (/health)",
        'qa': "Q1: Are there any skipped or mocked tests in the 80 passing tests?\nA1: No tests are skipped. All 80 tests execute live against local database fixtures and cryptographic primitives.\n\nQ2: How do you test credential tampering?\nA2: In test_vc.py, we mutate a single character in the credential subject after issuance and assert that verify_credential raises a cryptographic verification error.",
        'caveat': "Automated tests execute in Python 3.12 virtual environment under pytest 8.4."
    })

    # ==========================================================================
    # SLIDE 17: Transparent Limitations, Threats to Validity & Roadmap
    # ==========================================================================
    slide17 = create_base_slide("CRITICAL ANALYSIS \u00B7 HONEST BOUNDARIES & FUTURE SCOPE", "Transparent Limitations, Threats to Validity & Research Roadmap", 17)

    # Left Column: 7 Frank Limitations
    c17_l = add_card(slide17, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "7 Frank Technical Limitations & Threats to Validity")
    tf17_l = add_textbox(slide17, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    limitations = [
        ("1. Local Relational Storage", "Uses SQLite WAL rather than a distributed Byzantine consensus ledger (Ethereum/Hedera)."),
        ("2. Deterministic Telemetry Emulation", "Evaluated via LocalAgentConnector; wild LLM reasoning variance remains unmeasured."),
        ("3. Sensor Authenticity Assumption", "Assumes local gateway sensors are uncompromised; internal sensor spoofing is out of scope."),
        ("4. Static Hyperparameter Weights", "EMA weights (0.7/0.3) and matching weights (0.4/0.3/0.2/0.1) are fixed rather than learned."),
        ("5. Sequential Pipeline Boundary", "Multi-agent workflows run sequentially; parallel DAG fork-join execution is future work."),
        ("6. Closed-World Evaluation Scale", "Benchmarked with 8-10 agents; million-agent scale requires distributed DHT indexing."),
        ("7. False Positive Anomaly Risk", "Transient network hiccups can cause brief latency penalties, temporarily nudging trust into Suspicious.")
    ]
    for tag, desc in limitations:
        add_para(tf17_l, f"\u2022 {tag}:", size=9.5, bold=True, color=COLOR_AMBER)
        add_para(tf17_l, f"  {desc}", size=9, color=COLOR_SLATE, space_before=1)

    # Right Column: Future Roadmap
    c17_r = add_card(slide17, Inches(6.883), Inches(1.65), Inches(5.65), Inches(5.25), "Strategic Future Research Roadmap & Scalability Horizon")
    tf17_r = add_textbox(slide17, Inches(7.133), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf17_r, "PHASED RESEARCH & EVOLUTION ROADMAP", size=11, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf17_r, "TrustChain-AI establishes a strong foundation. Four key enhancements define the future research trajectory:", size=10, color=COLOR_SLATE, space_before=2)

    roadmap = [
        ("PHASE 1 \u00B7 DISTRIBUTED BYZANTINE VDR", "Integrate Hedera Hashgraph or Tendermint consensus to replace local SQLite with decentralized, multi-region Byzantine fault tolerance.", COLOR_CYAN),
        ("PHASE 2 \u00B7 ZERO-KNOWLEDGE TELEMETRY PROOFS", "Deploy zk-SNARKs so proprietary commercial AI agents can mathematically prove task compliance without disclosing proprietary model prompts or training weights.", COLOR_PURPLE),
        ("PHASE 3 \u00B7 REINFORCEMENT LEARNING WEIGHT ADAPTATION", "Replace static weighting constants (\u03b1=0.7, \u03b2=0.3) with context-aware reinforcement learning agents that tune weights based on environmental volatility.", COLOR_GREEN),
        ("PHASE 4 \u00B7 ASYNCHRONOUS DAG WORKFLOW ENGINE", "Upgrade sequential pipeline engine to support full DAG fork-join parallel execution with asynchronous synchronization barriers.", COLOR_AMBER)
    ]
    for p_tag, p_desc, p_col in roadmap:
        add_para(tf17_r, f"[{p_tag}]", size=9.5, bold=True, color=p_col, mono=True, space_before=6)
        add_para(tf17_r, p_desc, size=9, color=COLOR_SLATE, space_before=1)

    attach_notes(slide17, {
        'narrative': "Slide 17 is dedicated to scientific rigor, academic honesty, and future research. Every engineering project has explicit boundaries. Rather than overclaiming, we candidly acknowledge seven technical limitations: our registry is a local SQLite WAL instance rather than a Byzantine consensus ledger; agent telemetry was evaluated using deterministic simulation rather than live LLMs; gateway sensors are assumed trustworthy; mathematical weights are static rather than dynamically learned; workflows run sequentially rather than in parallel DAGs; testing was conducted in a closed-world registry of 8 to 10 agents; and transient network spikes can occasionally trigger brief false-positive latency penalties. On the right, we outline our four-phase future roadmap, including Byzantine ledger integration, zero-knowledge telemetry proofs, and reinforcement learning for dynamic weight tuning.",
        'takeaway': "Academic integrity requires candidly presenting technical boundaries and grounding future work in concrete engineering milestones.",
        'evidence': "docs/blackbook/TrustChain_AI_Blackbook_Part2.md (Chapter 6: Limitations and Future Work)",
        'qa': "Q1: What is the most significant limitation of TrustChain-AI today?\nA1: The reliance on a local SQLite VDR. While ideal for high-throughput single-node deployment, scaling across multi-tenant enterprise consortiums requires a distributed consensus ledger like Hedera or Tendermint.\n\nQ2: How could zk-SNARKs help autonomous agents?\nA2: Enterprise agents often process sensitive proprietary prompts. Zero-knowledge proofs would allow an agent to prove it complied with a security policy without revealing its proprietary input or output.",
        'caveat': "All identified limitations are directly reflected in Chapter 6 of the capstone blackbook."
    })

    # ==========================================================================
    # SLIDE 18: Capstone Defense Conclusion & Contributions
    # ==========================================================================
    slide18 = create_base_slide("PROJECT DEFENSE CONCLUSION \u00B7 KEY CONTRIBUTIONS", "Capstone Defense Conclusion, Core Contributions & Open Defense", 18)

    # Left Column: Three Core Contributions
    c18_l = add_card(slide18, Inches(0.8), Inches(1.65), Inches(5.65), Inches(5.25), "Summary of Three Core Research & Engineering Contributions")
    tf18_l = add_textbox(slide18, Inches(1.05), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf18_l, "1. Cryptographic Trust Anchor (W3C did:key & VCs)", size=11, bold=True, color=COLOR_CYAN)
    add_para(tf18_l, "Engineered decentralized agent identity using Ed25519 keypairs and multicodec encoding without central CA reliance. Implemented W3C Verifiable Credentials with URDNA2015 canonicalization and Ed25519Signature2020 proofs.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf18_l, "2. Dynamic Behavioral Telemetry & Adaptive Zero-Trust Gating", size=11, bold=True, color=COLOR_GREEN, space_before=8)
    add_para(tf18_l, "Replaced static binary authentication with continuous Exponential Moving Average trust scoring (T = 0.7\u00B7T + 0.3\u00B7P). Coupled trust scores directly to operation-level adaptive gating (blocking privileged actions in Suspicious, total quarantine below 0.50).", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf18_l, "3. Verifiable Traceability & Explainable Matching", size=11, bold=True, color=COLOR_PURPLE, space_before=8)
    add_para(tf18_l, "Built a 13-table SQLite WAL Verifiable Data Registry with quad-dimensional auditability. Created an explainable multi-factor candidate matching engine (MatchScore = 0.4T + 0.3(r/5) + 0.2\u03c3 + 0.1(1-\u03bb)) backed by a modern React 19 SOC portal.", size=9.5, color=COLOR_SLATE, space_before=2)

    add_para(tf18_l, "Empirical Validation: 80/80 passing pytest tests in 19.10s and successful Vite production build.", size=10, bold=True, color=COLOR_WHITE, space_before=8)

    # Right Column: Closing & Open Defense
    c18_r = add_card(slide18, Inches(6.883), Inches(1.65), Inches(5.65), Inches(5.25), "Capstone Defense Conclusion & Committee Q&A", border_color=COLOR_CYAN)
    tf18_r = add_textbox(slide18, Inches(7.133), Inches(2.2), Inches(5.15), Inches(4.5))

    add_para(tf18_r, "PROJECT DEFENSE SUMMARY", size=11, bold=True, color=COLOR_CYAN, mono=True)
    add_para(tf18_r, "TrustChain-AI delivers a complete, verifiable, closed-loop zero-trust architecture for autonomous AI agents. By replacing fragile static authentication with continuous behavioral evaluation, TrustChain-AI guarantees that autonomous agent collaboration remains secure, compliant, and accountable even in the presence of runtime model degradation or adversarial drift.", size=10, color=COLOR_WHITE, space_before=4)

    add_para(tf18_r, "ACADEMIC ACKNOWLEDGMENTS", size=10.5, bold=True, color=COLOR_AMBER, mono=True, space_before=10)
    add_para(tf18_r, "\u2022 Department of Computer Engineering\n\u2022 Mukesh Patel School of Technology Management & Engineering (MPSTME)\n\u2022 SVKM's NMIMS (Deemed-to-be University), Mumbai\n\u2022 Project Guide & Mentors for technical guidance throughout 2026\u20132027.", size=9.5, color=COLOR_SLATE, space_before=2)

    # Glowing Q&A Banner
    qa_card = slide18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.133), Inches(5.5), Inches(5.15), Inches(1.15))
    qa_card.fill.solid()
    qa_card.fill.fore_color.rgb = COLOR_PANEL_ALT
    qa_card.line.color.rgb = COLOR_CYAN
    qa_card.line.width = Pt(1.5)
    qatf = qa_card.text_frame
    qatf.word_wrap = True
    qatf.margin_left = qatf.margin_top = qatf.margin_right = qatf.margin_bottom = 0
    qap = qatf.paragraphs[0]
    qap.text = "DEFENSE COMMITTEE Q&A SESSION"
    qap.font.name = FONT_TITLE
    qap.font.size = Pt(13)
    qap.font.bold = True
    qap.font.color.rgb = COLOR_CYAN
    qap.alignment = PP_ALIGN.CENTER
    qap2 = qatf.add_paragraph()
    qap2.text = "Thank you for your time and evaluation.\nWe welcome questions from the examination committee."
    qap2.font.name = FONT_BODY
    qap2.font.size = Pt(10)
    qap2.font.color.rgb = COLOR_WHITE
    qap2.alignment = PP_ALIGN.CENTER
    qap2.space_before = Pt(3)

    attach_notes(slide18, {
        'narrative': "To conclude our defense, TrustChain-AI successfully solves the point-in-time security failure in autonomous multi-agent systems. We have demonstrated three core contributions: a cryptographic trust anchor built on W3C did:key and Verifiable Credentials; dynamic behavioral telemetry coupled with an Exponential Moving Average trust engine and adaptive policy gating; and a high-performance 13-table SQLite WAL Verifiable Data Registry with quad-dimensional auditability and an explainable matching engine. All claims are backed by 80 automated unit tests passing in 19.10 seconds and a verified production frontend. We express our deepest gratitude to the Department of Computer Engineering at MPSTME, SVKM's NMIMS, our project guide, and the examination committee. We are now ready for your questions.",
        'takeaway': "TrustChain-AI replaces static point-in-time authentication with a closed-loop, verifiable zero-trust architecture backed by 80 automated passing tests.",
        'evidence': "docs/blackbook/TrustChain_AI_Final_Blackbook.docx, tests/, akta/, frontend/",
        'qa': "Q1: If you had three more months, what would you prioritize next?\nA1: We would prioritize two things: first, deploying our VDR interface onto a Byzantine consensus network like Hedera Hashgraph for decentralized multi-organization consortiums; second, implementing zk-SNARK telemetry proofs so commercial LLM agents can prove compliance without leaking proprietary model prompts.\n\nQ2: What is the single biggest lesson learned from this capstone?\nA2: That cryptographic identity is merely the baseline of security. In autonomous agent systems, trust is not an identity attribute\u2014it is a dynamic, continuously earned behavioral property.",
        'caveat': "Final defense slide concluded with complete technical evidence and verification."
    })

    print("Slides 13-18 built successfully.")
