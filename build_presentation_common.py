import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
output_dir = os.path.join(BASE_DIR, "docs", "presentation")
assets_dir = os.path.join(output_dir, "TrustChain_AI_Slide_Assets")
pptx_path = os.path.join(output_dir, "TrustChain_AI_Final_Presentation.pptx")
os.makedirs(output_dir, exist_ok=True)
os.makedirs(assets_dir, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
COLOR_BG = RGBColor(8, 14, 26)         # #080E1A Dark Navy
COLOR_PANEL = RGBColor(17, 28, 53)     # #111C35 Deep Container
COLOR_PANEL_ALT = RGBColor(12, 20, 38) # #0C1426 Darker Container
COLOR_BORDER = RGBColor(30, 58, 110)   # #1E3A6E Subtle Border
COLOR_CYAN = RGBColor(0, 240, 255)     # #00F0FF Neon Cyan
COLOR_GREEN = RGBColor(0, 230, 118)    # #00E676 Emerald
COLOR_AMBER = RGBColor(255, 179, 0)    # #FFB300 Amber
COLOR_RED = RGBColor(255, 61, 113)     # #FF3D71 Rose Red
COLOR_PURPLE = RGBColor(157, 78, 221)  # #9D4EDD Purple
COLOR_WHITE = RGBColor(248, 250, 252)  # #F8FAFC Crisp White
COLOR_SLATE = RGBColor(148, 163, 184)  # #94A3B8 Cool Slate
COLOR_MUTED = RGBColor(100, 116, 139)  # #64748B Muted

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"

def create_base_slide(tracker_text, slide_title, slide_num):
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.733), Inches(1.05))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = tracker_text.upper()
    p0.font.name = FONT_MONO
    p0.font.size = Pt(9)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_CYAN
    
    p1 = tf.add_paragraph()
    p1.text = slide_title
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(21)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_before = Pt(3)
    
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_BORDER
    line.line.fill.background()
    
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.28))
    ftf = footer_box.text_frame
    ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
    fp = ftf.paragraphs[0]
    fp.text = f"TRUSTCHAIN-AI · B.TECH CAPSTONE DEFENSE · ACADEMIC YEAR 2026–2027                                                                               SLIDE {slide_num} OF 18"
    fp.font.name = FONT_MONO
    fp.font.size = Pt(7.5)
    fp.font.color.rgb = COLOR_MUTED
    
    return slide

def add_card(slide, left, top, width, height, title=None, border_color=COLOR_BORDER, fill_color=COLOR_PANEL):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    if title:
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), Inches(0.38))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
    return card

def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    return tf

def add_para(tf, text, size=10.5, bold=False, color=COLOR_WHITE, space_before=3, mono=False, align=PP_ALIGN.LEFT):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_MONO if mono else FONT_BODY
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.space_before = Pt(space_before)
    p.alignment = align
    return p

def attach_notes(slide, notes_dict):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "=== 1. VERBAL PRESENTATION NARRATIVE (30-60 SECONDS) ==="
    p0 = text_frame.paragraphs[0]
    p0.font.bold = True
    p0.font.size = Pt(10)
    
    p1 = text_frame.add_paragraph()
    p1.text = notes_dict.get('narrative', '')
    p1.space_before = Pt(3)
    
    p2 = text_frame.add_paragraph()
    p2.text = "\n=== 2. CORE TECHNICAL TAKEAWAY ==="
    p2.font.bold = True
    p2.space_before = Pt(5)
    p3 = text_frame.add_paragraph()
    p3.text = notes_dict.get('takeaway', '')
    
    p4 = text_frame.add_paragraph()
    p4.text = "\n=== 3. GROUNDING IMPLEMENTATION ARTIFACTS & CODE REFS ==="
    p4.font.bold = True
    p4.space_before = Pt(5)
    p5 = text_frame.add_paragraph()
    p5.text = notes_dict.get('evidence', '')
    
    p6 = text_frame.add_paragraph()
    p6.text = "\n=== 4. LIKELY EXAMINER QUESTIONS & AUTHORITATIVE DEFENSE ==="
    p6.font.bold = True
    p6.space_before = Pt(5)
    p7 = text_frame.add_paragraph()
    p7.text = notes_dict.get('qa', '')
    
    p8 = text_frame.add_paragraph()
    p8.text = "\n=== 5. SCIENTIFIC HONESTY & TECHNICAL CAVEATS ==="
    p8.font.bold = True
    p8.space_before = Pt(5)
    p9 = text_frame.add_paragraph()
    p9.text = notes_dict.get('caveat', '')

print("Common layout primitives verified.")
